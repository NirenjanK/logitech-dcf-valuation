"""Builds Logitech_DCF_Model.xlsx with live formulas. Source data: Logitech SEC filings (see README)."""
import sys, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.utils import get_column_letter as L

SCEN = int(sys.argv[1]) if len(sys.argv)>1 else 2
OUT  = sys.argv[2] if len(sys.argv)>2 else "Logitech_DCF_Model.xlsx"

BLUE=Font(name="Arial",size=10,color="0000FF"); BLK=Font(name="Arial",size=10)
BOLD=Font(name="Arial",size=10,bold=True); HDR=Font(name="Arial",size=10,bold=True,color="FFFFFF")
TITLE=Font(name="Arial",size=14,bold=True); NOTE=Font(name="Arial",size=9,italic=True,color="595959")
FILL=PatternFill("solid",fgColor="1F3864"); SEC=PatternFill("solid",fgColor="D9E1F2"); KEY=PatternFill("solid",fgColor="FFF2CC")
thin=Side(style="thin",color="808080")
M='#,##0.0;(#,##0.0);"-"'; K='#,##0;(#,##0);"-"'; P='0.0%'; X='0.0"x"'; D='$#,##0.00'

wb=Workbook()
def sheet(name,first=False):
    ws=wb.active if first else wb.create_sheet(); ws.title=name; ws.sheet_view.showGridLines=False; return ws
def put(ws,ref,v,font=BLK,fmt=None,fill=None,al=None):
    c=ws[ref]; c.value=v; c.font=font
    if fmt: c.number_format=fmt
    if fill: c.fill=fill
    if al: c.alignment=al
    return c
def header(ws,row,labels,c0=1):
    for i,t in enumerate(labels):
        c=ws.cell(row=row,column=c0+i,value=t); c.font=HDR; c.fill=FILL; c.alignment=Alignment(horizontal="center" if i else "left")
def section(ws,row,text,ncols=11):
    for i in range(1,ncols+1): ws.cell(row=row,column=i).fill=SEC
    ws.cell(row=row,column=1,value=text).font=BOLD

# ---------------- Historicals (USD thousands, as printed) ----------------
H=sheet("Historicals"); 
yrs=["FY2022","FY2023","FY2024","FY2025","FY2026"]
put(H,"A1","Logitech International S.A. – Historical financials (USD thousands, fiscal years ended 31 March)",TITLE)
put(H,"A2","Blue = hard-coded from SEC filings (Form 8-K Ex. 99.1 full-year results releases). Black = formula. Sources listed at bottom.",NOTE)
header(H,4,["USD thousands"]+yrs)
data = [
("sec","INCOME STATEMENT",None),
("sales","Net sales",[5481101,4538818,4298467,4554900,4840761]),
("cogs","Cost of goods sold",[3204072,2806438,2509418,2582745,2742407]),
("cogs_am","Amortization of intangibles (in cost of sales)",[14023,12865,11028,9554,7017]),
("gp","Gross profit (as reported)",[2263006,1719515,1778021,1962601,2091337]),
("ms","Marketing and selling",[1025899,809182,730310,814414,816604]),
("rd","Research and development",[291844,280796,287243,309008,316221]),
("ga","General and administrative",[148648,124652,155056,164014,167160]),
("am","Amortization of intangibles & acquisition-related costs",[16947,11843,10934,10695,6298]),
("imp","Impairment of intangible assets",[7000,0,3526,0,0]),
("cc","Change in fair value of contingent consideration",[-3509,0,-250,0,0]),
("rs","Restructuring charges, net",[2165,34573,3866,9615,9860]),
("opex","Total operating expenses (as reported)",[1488994,1261046,1190685,1307746,1316143]),
("ebit","Operating income (as reported)",[774012,458469,587336,654855,775194]),
("intinc","Interest income",[1246,18331,50636,54997,48246]),
("oth","Other income (expense), net  [FY22 derived: not printed]",[560,-13278,-16376,-2980,3079]),
("ebt","Income before income taxes (as reported)",[775818,463522,621596,706872,826519]),
("tax","Provision for income taxes",[131305,98947,9453,75343,115332]),
("ni","Net income (as reported)",[644513,364575,612143,631529,711187]),
("eps","Diluted EPS (USD)",[3.78,2.23,3.87,4.13,4.80]),
("dsh","Diluted weighted average shares (thousands)",[170414,163704,158171,152784,148208]),
("sec","BALANCE SHEET (31 March)",None),
("cash","Cash and cash equivalents",[1328716,1149023,1520842,1503205,1741546]),
("ar","Accounts receivable, net",[675604,630382,541715,454546,505867]),
("inv","Inventories",[933124,682893,422513,503747,489948]),
("oca","Other current assets",[135478,142876,146270,131211,177895]),
("tca","Total current assets (as reported)",[3072922,2605174,2631340,2592709,2915256]),
("ppe","Property, plant and equipment, net",[109807,121503,116589,113858,116454]),
("gw","Goodwill",[448175,454610,461978,463230,465417]),
("intang","Other intangible assets, net",[83779,63173,44603,24630,12386]),
("oa","Other assets",[320722,316293,350194,344077,339075]),
("ta","Total assets (as reported)",[4035405,3560753,3604704,3538504,3848588]),
("ap","Accounts payable",[636306,406968,448627,414586,530983]),
("acc","Accrued and other current liabilities",[784848,643139,637262,686503,781990]),
("tcl","Total current liabilities (as reported)",[1421154,1050107,1085889,1101089,1312973]),
("itp","Income taxes payable (non-current)",[83380,106391,112572,88483,86322]),
("oncl","Other non-current liabilities",[132133,146695,172590,221512,237899]),
("tl","Total liabilities (as reported)",[1636667,1303193,1371051,1411084,1637194]),
("eq","Total shareholders' equity (as reported)",[2398738,2257560,2233653,2127420,2211394]),
("tle","Total liabilities and shareholders' equity (as reported)",[4035405,3560753,3604704,3538504,3848588]),
("iss","Issued shares (thousands)",[173106,173106,173106,168994,160784]),
("tre","Treasury shares (thousands)",[7855,13763,19243,20485,17282]),
("sec","CASH FLOW STATEMENT",None),
("cni","Net income",[644513,364575,612143,631529,711187]),
("dep","Depreciation",[88361,76309,63065,59664,64139]),
("cam","Amortization of intangible assets",[30179,24407,21681,20098,13315]),
("sbc","Share-based compensation expense",[93479,70782,82889,89913,112392]),
("cfo","Net cash provided by operating activities",[298318,534010,1145116,842557,1037207]),
("capex","Purchases of property, plant and equipment",[89152,92353,55897,56128,61562]),
("bb","Share repurchases",[412022,418346,504203,588838,534939]),
("div","Dividends paid",[159410,158680,182305,207853,233059]),
("sec","NET SALES BY PRODUCT CATEGORY",None),
("c_gam","Gaming (incl. Streamlabs)",[None,1288313,1231063,1338467,1414206]),
("c_key","Keyboards & Combos",[None,836432,821441,882643,937551]),
("c_poi","Pointing Devices",[None,728357,742987,788784,858904]),
("c_vid","Video Collaboration",[None,677923,609361,626000,689040]),
("c_web","Webcams",[None,378688,325225,315520,326172]),
("c_tab","Tablet Accessories",[None,254374,254060,299540,336189]),
("c_hea","Headsets",[None,176576,168478,179710,179825]),
("c_oth","Other (mainly mobile and PC speakers)",[None,198155,145852,124236,98874]),
]
HR={}; r=5
for key,label,vals in data:
    if key=="sec":
        r+=1; section(H,r,label,6); r+=1; continue
    put(H,f"A{r}",label,BOLD if "reported" in label else BLK); HR[key]=r
    for i,v in enumerate(vals):
        if v is None: continue
        put(H,f"{L(2+i)}{r}",v,BLUE,"0.00" if key=="eps" else K)
    r+=1
r+=1
for t in ["Sources: Logitech Form 8-K Exhibit 99.1 – Q4 & full fiscal year results: FY2026 (sec.gov/Archives/edgar/data/0001032975/000103297526000014/exhibit991q4fy26.htm),",
          "FY2024 (…/000103297524000015/exhibit991q4fy24.htm) for FY2024/FY2023, FY2022 (…/000103297522000008/exhibit991q4fy22ng.htm) for FY2022.",
          "Preliminary full-year results releases; audited Form 10-K figures may differ marginally. FY2023 category data reflect the company's later reclassification.",
          "Category sales shown for FY2023–FY2026 only (FY2022 categories not on the same basis)."]:
    put(H,f"A{r}",t,NOTE); r+=1
H.column_dimensions["A"].width=58
for c in "BCDEF": H.column_dimensions[c].width=14
H.freeze_panes="B5"
def h(key,col): return f"Historicals!{col}{HR[key]}"

# ---------------- Inputs ----------------
I=sheet("Inputs")
put(I,"A1","Inputs & assumptions (blue = editable input)",TITLE)
put(I,"A3","Scenario selector (1 = Bear, 2 = Base, 3 = Bull)",BOLD); put(I,"B3",SCEN,BLUE,fill=KEY)
put(I,"A4","Active scenario"); put(I,"B4","=INDEX(A7:A9,B3)",BOLD)
header(I,6,["Scenario","Revenue growth Δ (pp per year, all categories)","Gross margin Δ (pp)"])
for i,(n,g,m) in enumerate([("Bear",-0.02,-0.01),("Base",0,0),("Bull",0.02,0.01)]):
    put(I,f"A{7+i}",n); put(I,f"B{7+i}",g,BLUE,P); put(I,f"C{7+i}",m,BLUE,P)
put(I,"A11","Active growth Δ"); put(I,"B11","=INDEX(B7:B9,B3)",BLK,P)
put(I,"A12","Active gross margin Δ"); put(I,"B12","=INDEX(C7:C9,B3)",BLK,P)

section(I,13,"Market data & capital structure",11)
rows=[(14,"Valuation date",dt.date(2026,10,17),"yyyy-mm-dd","Model date"),
 (15,"Share price (USD)",104.41,D,"LOGI (NASDAQ) close 16 Oct 2026 – stockanalysis.com"),
 (16,"Shares outstanding (m)",f"=({h('iss','F')}-{h('tre','F')})/1000","0.000","Issued less treasury, 31 Mar 2026 (excludes unvested RSUs; SBC is treated as a cash-like cost, not added back)"),
 (17,"Cash & equivalents (USD m)",f"={h('cash','F')}/1000",M,"31 Mar 2026 balance sheet"),
 (18,"Debt (USD m)",0,M,"No borrowings on the balance sheet; operating leases excluded")]
for rr,lab,v,fmt,src in rows:
    put(I,f"A{rr}",lab); put(I,f"B{rr}",v,BLK if isinstance(v,str) else BLUE,fmt); put(I,f"C{rr}",src,NOTE)
section(I,20,"Cost of capital (CAPM / WACC)",11)
w=[(21,"Risk-free rate (US 10-yr Treasury)",0.051,P,"FRED DGS10, 15 Oct 2026: 5.10%"),
 (22,"Raw equity beta",0.59,"0.00","stockanalysis.com (LOGI), Oct 2026"),
 (23,"Adjusted beta (Blume: 0.67 × raw + 0.33)","=0.67*B22+0.33","0.00","Pulls raw beta toward 1.0"),
 (24,"Equity risk premium",0.045,P,"Assumption – within the commonly used 4–5.5% range; test in Sensitivity"),
 (25,"Cost of equity","=B21+B23*B24",P,"Rf + β × ERP"),
 (26,"Pre-tax cost of debt",0.06,P,"Placeholder – inactive at 0% debt weight"),
 (27,"Tax rate for debt shield",0.15,P,""),
 (28,"Debt / (Debt + Equity)",0,P,"Logitech carries no financial debt"),
 (29,"WACC","=(1-B28)*B25+B28*B26*(1-B27)",P,"")]
for rr,lab,v,fmt,src in w:
    put(I,f"A{rr}",lab,BOLD if rr in (25,29) else BLK); put(I,f"B{rr}",v,BLK if isinstance(v,str) else BLUE,fmt,KEY if rr==29 else None); put(I,f"C{rr}",src,NOTE)
section(I,30,"Terminal value",11)
put(I,"A31","Terminal growth rate"); put(I,"B31",0.025,BLUE,P); put(I,"C31","Assumption – long-run nominal growth below the risk-free rate",NOTE)
put(I,"A32","Exit EV/EBITDA multiple"); put(I,"B32",14.0,BLUE,X); put(I,"C32","Assumption – below the c.15.5x implied by current price on FY2026 EBITDA",NOTE)

section(I,34,"Operating forecast drivers (Base case; scenario deltas apply on top)",11)
header(I,35,["Driver","","","","","","FY2027E","FY2028E","FY2029E","FY2030E","FY2031E"])
cats=[("c_gam","Gaming growth",[.05,.05,.04,.04,.03]),("c_key","Keyboards & Combos growth",[.04,.04,.035,.03,.03]),
      ("c_poi","Pointing Devices growth",[.06,.05,.04,.04,.03]),("c_vid","Video Collaboration growth",[.07,.06,.05,.04,.03]),
      ("c_web","Webcams growth",[.03,.03,.03,.025,.025]),("c_tab","Tablet Accessories growth",[.08,.06,.05,.04,.03]),
      ("c_hea","Headsets growth",[.02,.02,.02,.02,.02]),("c_oth","Other growth",[-.10,-.05,0,0,0])]
IR={}
rr=36
for key,lab,vals in cats:
    put(I,f"A{rr}",lab); IR[key]=rr
    for i,v in enumerate(vals): put(I,f"{L(7+i)}{rr}",v,BLUE,P)
    rr+=1
others=[("gm","Gross margin (after amortization in cost of sales)",[.432,.433,.434,.434,.434]),
 ("ms","Marketing & selling, % of sales",[.168,.167,.166,.165,.165]),
 ("rd","Research & development, % of sales",[.065]*5),
 ("ga","General & administrative, % of sales",[.034]*5),
 ("tax","Tax rate on EBIT",[.15]*5),
 ("dep","Depreciation, % of sales",[.013]*5),
 ("capex","Capex, % of sales",[.013]*5),
 ("nwc","Increase in NWC, % of incremental sales",[.05]*5)]
for key,lab,vals in others:
    put(I,f"A{rr}",lab); IR[key]=rr
    for i,v in enumerate(vals): put(I,f"{L(7+i)}{rr}",v,BLUE,P)
    rr+=1
put(I,f"A{rr+1}","Calibration notes: category growth tapers from FY2027 (Q1 FY27 guidance: sales +4–6% USD) toward 3%; gross margin held near FY2026's 43.2%; opex ratios from FY2026 actuals (S&M 16.9%, R&D 6.5%, G&A 3.5%);",NOTE)
put(I,f"A{rr+2}","tax 15% vs FY2026 effective 14.0%; depreciation and capex at FY2026 level (1.3% of sales). Share-based compensation stays inside operating costs (no add-back).",NOTE)
I.column_dimensions["A"].width=52; I.column_dimensions["B"].width=22; I.column_dimensions["C"].width=22
for c in "DEF": I.column_dimensions[c].width=3
for c in "GHIJK": I.column_dimensions[c].width=12

# ---------------- DCF ----------------
Dc=sheet("DCF")
put(Dc,"A1","Logitech – Discounted cash flow (USD millions unless stated)",TITLE)
put(Dc,"A2","=\"Scenario: \"&Inputs!B4&\"   |   WACC \"&TEXT(Inputs!B29,\"0.0%\")&\"   |   Terminal growth \"&TEXT(Inputs!B31,\"0.0%\")",NOTE)
cols=[L(i) for i in range(2,12)]  # B..K
header(Dc,4,["USD m"]+[f"FY{y}{'A' if y<=2026 else 'E'}" for y in range(2022,2032)])
put(Dc,"A5","Fiscal year end",NOTE)
for i,c in enumerate(cols): put(Dc,f"{c}5",dt.date(2022+i,3,31),NOTE,"dd-mmm-yy")
section(Dc,6,"Net sales by product category")
ckeys=[k for k,_,_ in cats]; cnames=["Gaming","Keyboards & Combos","Pointing Devices","Video Collaboration","Webcams","Tablet Accessories","Headsets","Other"]
for j,(ck,nm) in enumerate(zip(ckeys,cnames)):
    rr=7+j; put(Dc,f"A{rr}",nm)
    for c in "CDEF": put(Dc,f"{c}{rr}",f"={h(ck,c)}/1000",BLK,M)
    for c in "GHIJK":
        prev=L(ord(c)-64-1)
        put(Dc,f"{c}{rr}",f"={prev}{rr}*(1+Inputs!{c}{IR[ck]}+Inputs!$B$11)",BLK,M)
put(Dc,"A15","Total net sales",BOLD)
for c in "BCDEF": put(Dc,f"{c}15",f"={h('sales',c)}/1000",BOLD,M)
for c in "GHIJK": put(Dc,f"{c}15",f"=SUM({c}7:{c}14)",BOLD,M)
put(Dc,"A16","  growth %")
for c in cols[1:]: put(Dc,f"{c}16",f"={c}15/{L(ord(c)-64-1)}15-1",NOTE,P)
section(Dc,17,"Operating profit")
put(Dc,"A18","Gross profit",BOLD); put(Dc,"A19","  gross margin %")
for c in "BCDEF": put(Dc,f"{c}18",f"={h('gp',c)}/1000",BOLD,M)
for c in "GHIJK": put(Dc,f"{c}18",f"={c}15*(Inputs!{c}{IR['gm']}+Inputs!$B$12)",BOLD,M)
for c in cols: put(Dc,f"{c}19",f"={c}18/{c}15",NOTE,P)
for rr,key,lab,ik in [(20,"ms","Marketing & selling","ms"),(21,"rd","Research & development","rd"),(22,"ga","General & administrative","ga")]:
    put(Dc,f"A{rr}",lab)
    for c in "BCDEF": put(Dc,f"{c}{rr}",f"=-{h(key,c)}/1000",BLK,M)
    for c in "GHIJK": put(Dc,f"{c}{rr}",f"=-{c}15*Inputs!{c}{IR[ik]}",BLK,M)
put(Dc,"A23","Other operating items (amortization, impairment, restructuring)")
for c in "BCDEF": put(Dc,f"{c}23",f"=-({h('am',c)}+{h('imp',c)}+{h('cc',c)}+{h('rs',c)})/1000",BLK,M)
for c in "GHIJK": put(Dc,f"{c}23",0,BLUE,M)
put(Dc,"A24","EBIT (operating income)",BOLD)
for c in cols: put(Dc,f"{c}24",f"=SUM({c}18,{c}20:{c}23)",BOLD,M)
put(Dc,"A25","  EBIT margin %")
for c in cols: put(Dc,f"{c}25",f"={c}24/{c}15",NOTE,P)
put(Dc,"A26","Depreciation")
for c in "BCDEF": put(Dc,f"{c}26",f"={h('dep',c)}/1000",BLK,M)
for c in "GHIJK": put(Dc,f"{c}26",f"={c}15*Inputs!{c}{IR['dep']}",BLK,M)
put(Dc,"A27","Amortization of intangibles")
for c in "BCDEF": put(Dc,f"{c}27",f"={h('cam',c)}/1000",BLK,M)
for c in "GHIJK": put(Dc,f"{c}27",0,BLUE,M)
put(Dc,"A28","EBITDA",BOLD)
for c in cols: put(Dc,f"{c}28",f"={c}24+{c}26+{c}27",BOLD,M)
section(Dc,30,"Unlevered free cash flow")
put(Dc,"A31","Less: tax on EBIT"); put(Dc,"A32","NOPAT",BOLD); put(Dc,"A33","Plus: depreciation & amortization"); put(Dc,"A34","Less: capital expenditure")
put(Dc,"A35","Less: increase in net working capital"); put(Dc,"A36","Unlevered free cash flow (UFCF)",BOLD); put(Dc,"A37","  UFCF margin %")
put(Dc,"A38","Memo: reported FCF (operating cash flow – capex)",NOTE)
for c in "BCDEF": put(Dc,f"{c}38",f"=({h('cfo',c)}-{h('capex',c)})/1000",NOTE,M)
for c in "GHIJK":
    p=L(ord(c)-64-1)
    put(Dc,f"{c}31",f"=-{c}24*Inputs!{c}{IR['tax']}",BLK,M)
    put(Dc,f"{c}32",f"={c}24+{c}31",BOLD,M)
    put(Dc,f"{c}33",f"={c}26+{c}27",BLK,M)
    put(Dc,f"{c}34",f"=-{c}15*Inputs!{c}{IR['capex']}",BLK,M)
    put(Dc,f"{c}35",f"=-({c}15-{p}15)*Inputs!{c}{IR['nwc']}",BLK,M)
    put(Dc,f"{c}36",f"=SUM({c}32:{c}35)",BOLD,M)
    put(Dc,f"{c}37",f"={c}36/{c}15",NOTE,P)
section(Dc,39,"Discounting (mid-period convention, stub from valuation date)")
put(Dc,"A40","Fraction of year included"); put(Dc,"A41","Discount period (years, mid-period)"); put(Dc,"A42","Discount factor")
put(Dc,"A43","UFCF in period"); put(Dc,"A44","PV of UFCF",BOLD)
for c in "GHIJK":
    p=L(ord(c)-64-1)
    put(Dc,f"{c}40",f"=MIN(1,({c}5-Inputs!$B$14)/365)",BLK,"0.000")
    put(Dc,f"{c}41",f"={c}40/2" if c=="G" else f"=SUM($G$40:{p}40)+{c}40/2",BLK,"0.000")
    put(Dc,f"{c}42",f"=1/(1+Inputs!$B$29)^{c}41",BLK,"0.0000")
    put(Dc,f"{c}43",f"={c}36*{c}40",BLK,M)
    put(Dc,f"{c}44",f"={c}43*{c}42",BOLD,M)
section(Dc,46,"Valuation – perpetuity growth method (base)")
v=[(47,"Sum of PV of UFCF","=SUM(G44:K44)",M),(48,"Terminal-year UFCF (FY2031E)","=K36",M),
 (49,"Terminal value at end FY2031E","=B48*(1+Inputs!B31)/(Inputs!B29-Inputs!B31)",M),
 (50,"Years to end of FY2031E","=SUM(G40:K40)","0.000"),(51,"PV of terminal value","=B49/(1+Inputs!B29)^B50",M),
 (52,"Enterprise value","=B47+B51",M),(53,"  Terminal value % of EV","=B51/B52",P),
 (54,"Plus: cash","=Inputs!B17",M),(55,"Less: debt","=-Inputs!B18",M),(56,"Equity value","=SUM(B52,B54,B55)",M),
 (57,"Shares outstanding (m)","=Inputs!B16","0.000"),(58,"Implied value per share (USD)","=B56/B57",D),
 (59,"Current share price (USD)","=Inputs!B15",D),(60,"Upside / (downside) vs current","=B58/B59-1",P),
 (61,"Implied EV / FY2026 EBITDA","=B52/F28",X),(62,"Implied EV / FY2027E EBITDA","=B52/G28",X)]
for rr,lab,f,fmt in v:
    put(Dc,f"A{rr}",lab,BOLD if rr in (52,56,58) else BLK); put(Dc,f"B{rr}",f,BOLD if rr in (52,56,58) else BLK,fmt,KEY if rr==58 else None)
section(Dc,64,"Valuation – exit multiple method (cross-check)")
v2=[(65,"FY2031E EBITDA","=K28",M),(66,"Terminal value at exit multiple","=B65*Inputs!B32",M),(67,"PV of terminal value","=B66/(1+Inputs!B29)^B50",M),
 (68,"Enterprise value","=B47+B67",M),(69,"Equity value","=B68+B54+B55",M),(70,"Implied value per share (USD)","=B69/B57",D),
 (71,"Upside / (downside) vs current","=B70/B59-1",P),
 (72,"Perpetuity growth implied by exit multiple","=(B66*Inputs!B29-B48)/(B66+B48)",P),
 (73,"EV/EBITDA implied by perpetuity method","=B49/K28",X)]
for rr,lab,f,fmt in v2:
    put(Dc,f"A{rr}",lab,BOLD if rr in (68,70) else BLK); put(Dc,f"B{rr}",f,BOLD if rr in (68,70) else BLK,fmt,KEY if rr==70 else None)
Dc.column_dimensions["A"].width=54
for c in cols: Dc.column_dimensions[c].width=11.5
Dc.freeze_panes="B5"

# ---------------- Sensitivity ----------------
S=sheet("Sensitivity")
put(S,"A1","Sensitivity analysis – implied value per share (USD)",TITLE)
put(S,"A2","Every cell is a live formula (no Excel data tables), so the grids update with any change in Inputs.",NOTE)
def price_g(wc,gc):  # WACC cell, g cell
    return (f"=(SUMPRODUCT(DCF!$G$43:$K$43/(1+{wc})^DCF!$G$41:$K$41)"
            f"+DCF!$K$36*(1+{gc})/({wc}-{gc})/(1+{wc})^DCF!$B$50+Inputs!$B$17-Inputs!$B$18)/Inputs!$B$16")
def price_x(wc,xc):
    return (f"=(SUMPRODUCT(DCF!$G$43:$K$43/(1+{wc})^DCF!$G$41:$K$41)"
            f"+DCF!$K$28*{xc}/(1+{wc})^DCF!$B$50+Inputs!$B$17-Inputs!$B$18)/Inputs!$B$16")
put(S,"A4","Table 1: WACC (down) vs terminal growth (across) – perpetuity method",BOLD)
put(S,"A5","WACC \\ g",BOLD,al=Alignment(horizontal="center"))
for j,dg in enumerate([-0.01,-0.005,0,0.005,0.01]):
    put(S,f"{L(2+j)}5",f"=Inputs!$B$31+({dg})",BOLD,P,SEC)
for i,dw in enumerate([-0.01,-0.005,0,0.005,0.01]):
    rr=6+i; put(S,f"A{rr}",f"=Inputs!$B$29+({dw})",BOLD,P,SEC)
    for j in range(5):
        put(S,f"{L(2+j)}{rr}",price_g(f"$A{rr}",f"{L(2+j)}$5"),BLK,D,KEY if (i==2 and j==2) else None)
put(S,"A13","Table 2: WACC (down) vs exit EV/EBITDA multiple (across) – exit multiple method",BOLD)
put(S,"A14","WACC \\ multiple",BOLD,al=Alignment(horizontal="center"))
for j,dx in enumerate([-2,-1,0,1,2]): put(S,f"{L(2+j)}14",f"=Inputs!$B$32+({dx})",BOLD,X,SEC)
for i,dw in enumerate([-0.01,-0.005,0,0.005,0.01]):
    rr=15+i; put(S,f"A{rr}",f"=Inputs!$B$29+({dw})",BOLD,P,SEC)
    for j in range(5):
        put(S,f"{L(2+j)}{rr}",price_x(f"$A{rr}",f"{L(2+j)}$14"),BLK,D,KEY if (i==2 and j==2) else None)
put(S,"A22","Table 3: Upside / (downside) vs current price – Table 1 re-expressed",BOLD)
put(S,"A23","WACC \\ g",BOLD,al=Alignment(horizontal="center"))
for j in range(5): put(S,f"{L(2+j)}23",f"={L(2+j)}5",BOLD,P,SEC)
for i in range(5):
    rr=24+i; put(S,f"A{rr}",f"=A{6+i}",BOLD,P,SEC)
    for j in range(5): put(S,f"{L(2+j)}{rr}",f"={L(2+j)}{6+i}/Inputs!$B$15-1",BLK,P)
S.column_dimensions["A"].width=24
for c in "BCDEF": S.column_dimensions[c].width=13

# ---------------- Summary (+ football field) ----------------
Sm=wb.create_sheet("Summary",0); Sm.sheet_view.showGridLines=False
put(Sm,"A1","Logitech International (NASDAQ: LOGI) – DCF valuation summary",TITLE)
put(Sm,"A2","=\"Scenario: \"&Inputs!B4&\"   |   Valuation date \"&TEXT(Inputs!B14,\"d mmm yyyy\")&\"   |   USD\"",NOTE)
rows=[(4,"Implied value per share – perpetuity growth","=DCF!B58",D),(5,"Implied value per share – exit multiple","=DCF!B70",D),
 (6,"Current share price","=Inputs!B15",D),(7,"Upside / (downside) – perpetuity growth","=DCF!B60",P),
 (8,"Enterprise value (USD m)","=DCF!B52",M),(9,"Equity value (USD m)","=DCF!B56",M),(10,"WACC","=Inputs!B29",P),
 (11,"Terminal growth","=Inputs!B31",P),(12,"Terminal value % of EV","=DCF!B53",P),(13,"Model integrity checks","=Checks!B3",None)]
for rr,lab,f,fmt in rows:
    put(Sm,f"A{rr}",lab,BOLD if rr<=5 else BLK); put(Sm,f"B{rr}",f,BOLD,fmt,KEY if rr in (4,5) else None)
section(Sm,15,"Football field (USD per share)",5)
header(Sm,16,["Method","Low","High","Range"])
ff=[("DCF – perpetuity (WACC ±1%, g ±1%)","=MIN(Sensitivity!B6:F10)","=MAX(Sensitivity!B6:F10)"),
    ("DCF – exit multiple (WACC ±1%, ±2x)","=MIN(Sensitivity!B15:F19)","=MAX(Sensitivity!B15:F19)"),
    ("52-week trading range","83.32","129.66")]
for i,(n,lo,hi) in enumerate(ff):
    rr=17+i; put(Sm,f"A{rr}",n)
    put(Sm,f"B{rr}",lo if lo.startswith("=") else float(lo),BLK if lo.startswith("=") else BLUE,D)
    put(Sm,f"C{rr}",hi if hi.startswith("=") else float(hi),BLK if hi.startswith("=") else BLUE,D)
    put(Sm,f"D{rr}",f"=C{rr}-B{rr}",BLK,D)
put(Sm,"A20","52-week range: stockanalysis.com, 16 Oct 2026. Current price marker not charted; see B6.",NOTE)
ch=BarChart(); ch.type="bar"; ch.grouping="stacked"; ch.overlap=100; ch.title="Valuation range per share (USD)"
cats_ref=Reference(Sm,min_col=1,min_row=17,max_row=19)
s1=Reference(Sm,min_col=2,min_row=16,max_row=19); s2=Reference(Sm,min_col=4,min_row=16,max_row=19)
ch.add_data(s1,titles_from_data=True); ch.add_data(s2,titles_from_data=True); ch.set_categories(cats_ref)
ch.series[0].graphicalProperties.noFill=True; ch.series[0].graphicalProperties.line.noFill=True
ch.series[1].graphicalProperties.solidFill="1F3864"
ch.legend=None; ch.height=7; ch.width=18; ch.y_axis.delete=False; ch.x_axis.delete=False
Sm.add_chart(ch,"A22")
Sm.column_dimensions["A"].width=48
for c in "BCD": Sm.column_dimensions[c].width=14

# ---------------- Checks ----------------
Ck=sheet("Checks")
put(Ck,"A1","Model integrity checks (tolerance: 1 = USD 1 thousand rounding)",TITLE)
put(Ck,"A3","Overall status",BOLD)
checks=[]
def addcheck(label,expr_by_col):
    checks.append((label,expr_by_col))
hc=lambda k,c:h(k,c)
addcheck("Gross profit = sales − COGS − amortization in COGS",lambda c:f"={hc('sales',c)}-{hc('cogs',c)}-{hc('cogs_am',c)}-{hc('gp',c)}")
addcheck("Total opex = sum of opex lines",lambda c:f"={hc('ms',c)}+{hc('rd',c)}+{hc('ga',c)}+{hc('am',c)}+{hc('imp',c)}+{hc('cc',c)}+{hc('rs',c)}-{hc('opex',c)}")
addcheck("EBIT = gross profit − opex",lambda c:f"={hc('gp',c)}-{hc('opex',c)}-{hc('ebit',c)}")
addcheck("Pre-tax income = EBIT + interest + other",lambda c:f"={hc('ebit',c)}+{hc('intinc',c)}+{hc('oth',c)}-{hc('ebt',c)}")
addcheck("Net income = pre-tax income − tax",lambda c:f"={hc('ebt',c)}-{hc('tax',c)}-{hc('ni',c)}")
addcheck("Current assets = sum of components",lambda c:f"={hc('cash',c)}+{hc('ar',c)}+{hc('inv',c)}+{hc('oca',c)}-{hc('tca',c)}")
addcheck("Total assets = sum of components",lambda c:f"={hc('tca',c)}+{hc('ppe',c)}+{hc('gw',c)}+{hc('intang',c)}+{hc('oa',c)}-{hc('ta',c)}")
addcheck("Current liabilities = AP + accrued",lambda c:f"={hc('ap',c)}+{hc('acc',c)}-{hc('tcl',c)}")
addcheck("Total liabilities = current + non-current",lambda c:f"={hc('tcl',c)}+{hc('itp',c)}+{hc('oncl',c)}-{hc('tl',c)}")
addcheck("Assets = liabilities + equity",lambda c:f"={hc('ta',c)}-{hc('tl',c)}-{hc('eq',c)}")
addcheck("Cash-flow net income = income-statement net income",lambda c:f"={hc('cni',c)}-{hc('ni',c)}")
header(Ck,5,["Check (difference, USD thousands)"]+yrs+["Pass?"])
r=6
for lab,fn in checks:
    put(Ck,f"A{r}",lab)
    for i,c in enumerate("BCDEF"): put(Ck,f"{L(2+i)}{r}",fn(c),BLK,K)
    put(Ck,f"G{r}",f"=IF(MAX(ABS(MIN(B{r}:F{r})),ABS(MAX(B{r}:F{r})))<=1,\"OK\",\"CHECK\")",BOLD); r+=1
put(Ck,f"A{r}","Category sales = total net sales (FY2023–FY2026)")
for i,c in enumerate("CDEF"):
    put(Ck,f"{c}{r}","="+"+".join(h(k,c) for k in ckeys)+f"-{h('sales',c)}",BLK,K)
put(Ck,f"G{r}",f"=IF(MAX(ABS(MIN(C{r}:F{r})),ABS(MAX(C{r}:F{r})))<=1,\"OK\",\"CHECK\")",BOLD); r+=1
put(Ck,f"A{r}","Forecast category sum = total net sales (FY2027E–FY2031E)")
for c in "GHIJK": pass
put(Ck,f"B{r}","=SUM(DCF!G7:K14)-SUM(DCF!G15:K15)",BLK,M); put(Ck,f"G{r}",f"=IF(ABS(B{r})<0.001,\"OK\",\"CHECK\")",BOLD); r+=1
put(Ck,f"A{r}","WACC exceeds terminal growth"); put(Ck,f"B{r}","=Inputs!B29-Inputs!B31",BLK,P); put(Ck,f"G{r}",f"=IF(B{r}>0,\"OK\",\"CHECK\")",BOLD); r+=1
put(Ck,f"A{r}","Sensitivity base cell = DCF implied price"); put(Ck,f"B{r}","=Sensitivity!D8-DCF!B58",BLK,"0.0000"); put(Ck,f"G{r}",f"=IF(ABS(B{r})<0.001,\"OK\",\"CHECK\")",BOLD); r+=1
put(Ck,f"A{r}","Exit-multiple sensitivity base cell = DCF exit-multiple price"); put(Ck,f"B{r}","=Sensitivity!D17-DCF!B70",BLK,"0.0000"); put(Ck,f"G{r}",f"=IF(ABS(B{r})<0.001,\"OK\",\"CHECK\")",BOLD); r+=1
put(Ck,"B3",f"=IF(COUNTIF(G6:G{r-1},\"CHECK\")=0,\"ALL CHECKS PASS\",COUNTIF(G6:G{r-1},\"CHECK\")&\" CHECK(S) FAILING\")",BOLD,fill=KEY)
Ck.column_dimensions["A"].width=58
for c in "BCDEFG": Ck.column_dimensions[c].width=13
if "Sheet" in wb.sheetnames: del wb["Sheet"]
wb.save(OUT)
